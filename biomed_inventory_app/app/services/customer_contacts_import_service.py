from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app import erp_models as erp
from app.config.database import DATA_ROOT
from app.database import SessionLocal
from app.models.foundation import Manufacturer, Supplier
from app.models.foundation import AuditEvent, DataValidationError, ImportBatch, ImportRow

SOURCE = "outlook_email_history_import"
EXPECTED_SHEET = "Customer Contacts"
EXPECTED_COLUMNS = ["Name", "Email", "Company / Domain", "Emails Exchanged", "Role", "Phone", "Notes"]
EMAIL_RE = re.compile(r"^[A-Z0-9._%+\-']+@[A-Z0-9.\-]+\.[A-Z]{2,}$", re.I)
PUBLIC_DOMAINS = {"gmail.com", "hotmail.com", "outlook.com", "yahoo.com", "live.com", "icloud.com"}
SHARED_TERMS = {
    "purchasing": "purchasing",
    "procurement": "purchasing",
    "service": "service",
    "support": "technical_support",
    "sales": "sales",
    "info": "shared_inbox",
    "admin": "administration",
    "finance": "finance",
    "training": "training",
    "reception": "administration",
    "biomedical": "biomedical_engineering",
    "clinicalengineering": "clinical_engineering",
    "clinical.engineering": "clinical_engineering",
}
AUTOMATED_TERMS = {"noreply", "no-reply", "notification", "notifications", "mailer-daemon", "account-security", "learning.notifications", "automated", "system"}

CONTACT_IMPORT_MAX_FILE_SIZE_MB = int(os.getenv("CONTACT_IMPORT_MAX_FILE_SIZE_MB", "250"))
CONTACT_IMPORT_MAX_ROWS = int(os.getenv("CONTACT_IMPORT_MAX_ROWS", "1000000"))
CONTACT_IMPORT_BATCH_SIZE = int(os.getenv("CONTACT_IMPORT_BATCH_SIZE", "1000"))
CONTACT_IMPORT_DB_BATCH_SIZE = int(os.getenv("CONTACT_IMPORT_DB_BATCH_SIZE", "500"))
CONTACT_IMPORT_PREVIEW_ROWS = int(os.getenv("CONTACT_IMPORT_PREVIEW_ROWS", "100"))
CONTACT_IMPORT_ERROR_LIMIT = int(os.getenv("CONTACT_IMPORT_ERROR_LIMIT", "10000"))
CONTACT_IMPORT_WORKER_COUNT = int(os.getenv("CONTACT_IMPORT_WORKER_COUNT", "1"))
CONTACT_IMPORT_RETENTION_DAYS = int(os.getenv("CONTACT_IMPORT_RETENTION_DAYS", "30"))
CONTACT_IMPORT_STORAGE = Path(os.getenv("CONTACT_IMPORT_STORAGE", str(DATA_ROOT / "imports" / "customer_contacts")))
CONTACT_IMPORT_STORAGE.mkdir(parents=True, exist_ok=True)
ALLOWED_CONTACT_IMPORT_EXTENSIONS = {".xlsx", ".csv"}


def normalize_email(value: Any) -> str:
    return str(value or "").strip().casefold()


def safe_filename(value: str) -> str:
    stem = Path(value or "customer-contacts").stem
    return re.sub(r"[^A-Za-z0-9_.-]+", "-", stem).strip(".-")[:80] or "customer-contacts"


def email_domain(email: str) -> str:
    return email.split("@", 1)[1] if "@" in email else ""


def engagement_level(count: int) -> str:
    if count >= 30:
        return "high"
    if count >= 10:
        return "active"
    if count >= 3:
        return "moderate"
    return "low"


def parse_count(value: Any) -> int:
    try:
        return max(0, int(float(str(value or "0").replace(",", "").strip() or 0)))
    except Exception:
        return 0


def normalize_phone(value: Any) -> tuple[str | None, str | None]:
    original = str(value or "").strip()
    digits = re.sub(r"[^\d+]", "", original)
    return (digits or None, original or None)


def split_name(value: str) -> tuple[str | None, str | None]:
    cleaned = re.sub(r"\s+", " ", value or "").strip()
    if not cleaned or any(term in cleaned.casefold() for term in SHARED_TERMS):
        return None, None
    parts = cleaned.split(" ")
    if 2 <= len(parts) <= 3 and all(re.match(r"^[A-Za-z][A-Za-z'\-]+$", part) for part in parts):
        return parts[0], " ".join(parts[1:])
    return None, None


def detect_shared(display_name: str, email: str) -> tuple[bool, str]:
    text = f"{display_name} {email.split('@', 1)[0] if '@' in email else email}".casefold()
    for term, contact_type in SHARED_TERMS.items():
        if term in text:
            return True, contact_type
    return False, "person"


def detect_automated(display_name: str, email: str) -> bool:
    text = f"{display_name} {email.split('@', 1)[0] if '@' in email else email}".casefold()
    return any(term in text for term in AUTOMATED_TERMS)


def quality_status(normalized: dict[str, Any], messages: list[str]) -> str:
    if not normalized.get("organization_type") or normalized.get("organization_type") == "unknown":
        return "needs_organization_match"
    if normalized.get("is_automated_address") or normalized.get("is_shared_inbox") or messages:
        return "needs_review"
    if not normalized.get("role_title"):
        return "needs_role"
    if not normalized.get("phone"):
        return "needs_phone"
    return "complete"


class CustomerContactsImportService:
    def __init__(self, db: Session, user_id: int | None = None):
        self.db = db
        self.user_id = user_id

    def preview(self, content: bytes, filename: str, sheet_name: str | None = None) -> dict[str, Any]:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook[EXPECTED_SHEET] if EXPECTED_SHEET in workbook.sheetnames else workbook.active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        header_map = {header: index for index, header in enumerate(headers)}
        rows = []
        issues = []
        batch = ImportBatch(
            source_type="customer_contacts",
            source_filename=filename,
            source_checksum=hashlib.sha256(content).hexdigest(),
            imported_by_id=self.user_id,
            status="preview",
            total_rows=max(0, sheet.max_row - 1),
            notes=json.dumps({"worksheet": sheet.title, "columns": headers}),
        )
        self.db.add(batch)
        self.db.flush()
        for excel_row in sheet.iter_rows(min_row=2, values_only=True):
            raw = {column: excel_row[header_map[column]] if column in header_map and header_map[column] < len(excel_row) else "" for column in EXPECTED_COLUMNS}
            staged = self.stage_row(batch.id, len(rows) + 2, raw)
            rows.append(staged)
            for issue in staged["issues"]:
                if issue["severity"] == "error":
                    issues.append(issue)
        batch.processed_rows = len(rows)
        batch.successful_rows = sum(1 for row in rows if row["import_status"] in {"ready_create", "ready_update", "skipped"})
        batch.failed_rows = sum(1 for row in rows if row["import_status"] == "error")
        self._audit("customer_contacts_preview_generated", "import_batch", str(batch.id), {"filename": filename, "rows": len(rows)})
        self.db.commit()
        return self.import_payload(batch.id)

    def stage_row(self, batch_id: int, row_number: int, raw: dict[str, Any]) -> dict[str, Any]:
        name = str(raw.get("Name") or "").strip()
        imported_email = normalize_email(raw.get("Email"))
        issues: list[dict[str, Any]] = []
        if not imported_email:
            normalized = {"display_name": name, "import_action": "skip", "skip_reason": "domain_heading_or_blank_email"}
            import_row = ImportRow(import_batch_id=batch_id, row_number=row_number, raw_data=raw, normalized_data=normalized, processing_status="skipped", warning_message="Skipped grouping/domain row")
            self.db.add(import_row)
            self.db.flush()
            return self.row_payload(import_row, [])
        if not EMAIL_RE.match(imported_email):
            issues.append(self.issue(row_number, "email", imported_email, "invalid_email", "Email address is invalid", "error"))
        domain = email_domain(imported_email)
        supplied_domain = str(raw.get("Company / Domain") or "").strip().casefold()
        if supplied_domain and supplied_domain != domain and "@" not in supplied_domain:
            issues.append(self.issue(row_number, "company_domain", supplied_domain, "domain_mismatch", "Company / Domain differs from the email-derived domain", "warning"))
        shared, contact_type = detect_shared(name, imported_email)
        automated = detect_automated(name, imported_email)
        phone, phone_original = normalize_phone(raw.get("Phone"))
        first_name, last_name = split_name(name)
        org = self.match_organization(domain)
        count = parse_count(raw.get("Emails Exchanged"))
        existing = self.find_existing(imported_email)
        normalized = {
            "display_name": name,
            "first_name": first_name,
            "last_name": last_name,
            "email": imported_email,
            "normalized_email": imported_email,
            "email_domain": domain,
            "emails_exchanged": count,
            "engagement_level": engagement_level(count),
            "phone": phone,
            "phone_original": phone_original,
            "role_title": str(raw.get("Role") or "").strip() or None,
            "notes": str(raw.get("Notes") or "").strip() or None,
            "is_shared_inbox": shared,
            "is_automated_address": automated,
            "contact_type": "shared_inbox" if automated else contact_type,
            "organization_type": org.get("organization_type", "unknown"),
            "client_id": org.get("client_id"),
            "manufacturer_id": org.get("manufacturer_id"),
            "supplier_id": org.get("supplier_id"),
            "suggested_organization": org.get("name"),
            "existing_contact_id": existing.id if existing else None,
            "import_action": "update_existing_contact" if existing else "create_contact",
            "source": SOURCE,
            "source_reference": json.dumps(raw),
        }
        if not phone:
            issues.append(self.issue(row_number, "phone", "", "phone_missing", "Phone is missing; imported phone numbers require verification when present", "warning"))
        if phone:
            issues.append(self.issue(row_number, "phone", phone_original or phone, "phone_unverified", "Automatically extracted phone requires verification", "warning"))
        if not normalized["role_title"]:
            issues.append(self.issue(row_number, "role", "", "role_missing", "Role is missing and must be enriched manually", "warning"))
        if domain in PUBLIC_DOMAINS:
            normalized["organization_type"] = "personal"
            normalized["client_id"] = normalized["manufacturer_id"] = normalized["supplier_id"] = None
            issues.append(self.issue(row_number, "email_domain", domain, "public_email_domain", "Public email domains are not mapped to organizations automatically", "warning"))
        elif not org:
            issues.append(self.issue(row_number, "email_domain", domain, "organization_match_required", "Domain requires organization review", "warning"))
        normalized["data_quality_status"] = quality_status(normalized, [i["error_code"] for i in issues if i["severity"] == "warning"])
        status = "error" if any(i["severity"] == "error" for i in issues) else normalized["import_action"].replace("_existing_contact", "")
        import_row = ImportRow(import_batch_id=batch_id, row_number=row_number, raw_data=raw, normalized_data=normalized, processing_status=status, matched_client_id=normalized.get("client_id"), warning_message="; ".join(i["error_message"] for i in issues if i["severity"] == "warning") or None, error_message="; ".join(i["error_message"] for i in issues if i["severity"] == "error") or None)
        self.db.add(import_row)
        self.db.flush()
        for item in issues:
            values = dict(item)
            values.pop("row_number", None)
            self.db.add(DataValidationError(import_batch_id=batch_id, import_row_id=import_row.id, **values))
        return self.row_payload(import_row, issues)

    def confirm(self, import_id: int, decisions: list[dict[str, Any]] | None = None) -> dict[str, int]:
        decisions_by_row = {int(item["row_number"]): item for item in decisions or [] if item.get("row_number")}
        batch = self.db.get(ImportBatch, import_id)
        if not batch or batch.source_type != "customer_contacts":
            raise LookupError("Customer contacts import not found")
        created = updated = rejected = ignored = 0
        try:
            for row in self.db.query(ImportRow).filter_by(import_batch_id=import_id).order_by(ImportRow.row_number).all():
                data = dict(row.normalized_data or {})
                decision = decisions_by_row.get(row.row_number, {})
                action = decision.get("action") or data.get("import_action")
                if row.processing_status == "error" or action in {"reject", "rejected"}:
                    row.processing_status = "rejected"
                    rejected += 1
                    self._audit("customer_contact_rejected", "import_row", str(row.id), {"row_number": row.row_number, "reason": decision.get("reason")})
                    continue
                if action in {"ignore", "skip"} or row.processing_status == "skipped":
                    row.processing_status = "skipped"
                    ignored += 1
                    continue
                self.apply_decision(data, decision)
                contact, was_created = self.upsert_contact(data)
                row.processing_status = "imported"
                row.processed_at = datetime.now(UTC)
                if was_created:
                    created += 1
                    self._audit("customer_contact_created", "contact", str(contact.id), {"row_number": row.row_number, "contact_id": contact.id, "email": contact.email})
                else:
                    updated += 1
                    self._audit("customer_contact_updated", "contact", str(contact.id), {"row_number": row.row_number, "contact_id": contact.id, "email": contact.email})
            batch.status = "completed"
            batch.completed_at = datetime.now(UTC)
            batch.processed_rows = created + updated + rejected + ignored
            batch.successful_rows = created + updated
            batch.failed_rows = rejected
            self._audit("customer_contacts_import_confirmed", "import_batch", str(import_id), {"created": created, "updated": updated, "rejected": rejected, "ignored": ignored})
            self.db.commit()
        except Exception:
            self.db.rollback()
            raise
        return {"created": created, "updated": updated, "rejected": rejected, "ignored": ignored}

    def upsert_contact(self, data: dict[str, Any]) -> tuple[erp.Contact, bool]:
        contact = self.find_existing(data["normalized_email"])
        created = contact is None
        if created:
            contact = erp.Contact(name=data.get("display_name") or data["email"], email=data["email"], normalized_email=data["normalized_email"])
            self.db.add(contact)
        old_phone = contact.phone
        contact.display_name = contact.display_name or data.get("display_name")
        contact.name = contact.name or data.get("display_name") or data["email"]
        contact.first_name = contact.first_name or data.get("first_name")
        contact.last_name = contact.last_name or data.get("last_name")
        contact.email = data["email"]
        contact.normalized_email = data["normalized_email"]
        contact.email_domain = data.get("email_domain")
        contact.role_title = contact.role_title or data.get("role_title")
        contact.title = contact.title or data.get("role_title")
        if not contact.phone or (not contact.phone_verified and data.get("phone")):
            contact.phone = data.get("phone") or contact.phone
            contact.phone_original = data.get("phone_original") or contact.phone_original
        else:
            contact.phone = old_phone
        if int(data.get("emails_exchanged") or 0) > int(contact.emails_exchanged or 0):
            contact.emails_exchanged = data.get("emails_exchanged")
            contact.engagement_level = data.get("engagement_level")
        contact.contact_type = data.get("contact_type") or contact.contact_type
        contact.organization_type = data.get("organization_type") or contact.organization_type
        for field in ("client_id", "manufacturer_id", "supplier_id", "client_site_id"):
            if getattr(contact, field, None) is None and data.get(field):
                setattr(contact, field, data.get(field))
        contact.is_shared_inbox = bool(data.get("is_shared_inbox"))
        contact.is_automated_address = bool(data.get("is_automated_address"))
        contact.source = SOURCE
        contact.source_reference = data.get("source_reference")
        contact.data_quality_status = data.get("data_quality_status") or contact.data_quality_status
        contact.last_imported_at = datetime.now(UTC)
        if data.get("notes") and not contact.notes:
            contact.notes = data["notes"]
        self.db.flush()
        return contact, created

    def import_payload(self, import_id: int) -> dict[str, Any]:
        batch = self.db.get(ImportBatch, import_id)
        if not batch:
            raise LookupError("Import not found")
        rows = [self.row_payload(row) for row in self.db.query(ImportRow).filter_by(import_batch_id=import_id).order_by(ImportRow.row_number).all()]
        return {
            "import_id": batch.id,
            "filename": batch.source_filename,
            "status": batch.status,
            "total_rows": batch.total_rows,
            "processed_rows": batch.processed_rows,
            "successful_rows": batch.successful_rows,
            "failed_rows": batch.failed_rows,
            "rows": rows,
            "metrics": self.metrics(rows),
        }

    def report_csv(self, import_id: int) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Import status", "Name", "Email", "Domain", "Emails exchanged", "Engagement level", "Phone", "Role", "Suggested organization", "Organization type", "Shared inbox", "Automated address", "Data-quality status", "Validation messages", "Import action"])
        for row in self.import_payload(import_id)["rows"]:
            writer.writerow([row.get("import_status"), row.get("display_name"), row.get("email"), row.get("email_domain"), row.get("emails_exchanged"), row.get("engagement_level"), row.get("phone"), row.get("role_title"), row.get("suggested_organization"), row.get("organization_type"), row.get("is_shared_inbox"), row.get("is_automated_address"), row.get("data_quality_status"), "; ".join(row.get("validation_messages") or []), row.get("import_action")])
        return output.getvalue().encode("utf-8")

    def row_payload(self, row: ImportRow, issues: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        data = dict(row.normalized_data or {})
        if issues is None:
            issues = [
                {
                    "field_name": item.field_name,
                    "raw_value": item.raw_value,
                    "error_code": item.error_code,
                    "error_message": item.error_message,
                    "severity": item.severity,
                }
                for item in self.db.query(DataValidationError).filter_by(import_row_id=row.id).all()
            ]
        data.update({
            "import_row_id": row.id,
            "row_number": row.row_number,
            "import_status": row.processing_status,
            "validation_messages": [item.get("error_message") for item in issues],
            "issues": issues,
        })
        return data

    def metrics(self, rows: list[dict[str, Any]]) -> dict[str, int]:
        return {
            "contacts": sum(1 for row in rows if row.get("email")),
            "skipped_group_rows": sum(1 for row in rows if row.get("import_status") == "skipped"),
            "invalid_rows": sum(1 for row in rows if row.get("import_status") == "error"),
            "duplicates": sum(1 for row in rows if row.get("existing_contact_id")),
            "needs_organization_match": sum(1 for row in rows if row.get("data_quality_status") == "needs_organization_match"),
            "missing_role": sum(1 for row in rows if not row.get("role_title") and row.get("email")),
            "missing_phone": sum(1 for row in rows if not row.get("phone") and row.get("email")),
            "shared_inboxes": sum(1 for row in rows if row.get("is_shared_inbox")),
            "automated_addresses": sum(1 for row in rows if row.get("is_automated_address")),
            "high_engagement": sum(1 for row in rows if row.get("engagement_level") == "high"),
        }

    def match_organization(self, domain: str) -> dict[str, Any]:
        if not domain or domain in PUBLIC_DOMAINS:
            return {}
        mapping = self.db.query(erp.OrganizationDomainMapping).filter(func.lower(erp.OrganizationDomainMapping.domain) == domain, erp.OrganizationDomainMapping.status == "approved").first()
        if mapping:
            return self._org_payload(mapping.organization_type, mapping.organization_id)
        for client in self.db.query(erp.Client).all():
            if self._domain_matches_name(domain, client.name):
                return {"organization_type": "client", "client_id": client.id, "name": client.name}
        for manufacturer in self.db.query(Manufacturer).all():
            if self._domain_matches_name(domain, manufacturer.name) or domain in str(manufacturer.website or "").casefold():
                return {"organization_type": "manufacturer", "manufacturer_id": manufacturer.id, "name": manufacturer.name}
        for supplier in self.db.query(Supplier).all():
            if self._domain_matches_name(domain, supplier.name) or domain in str(supplier.website or "").casefold():
                return {"organization_type": "supplier", "supplier_id": supplier.id, "name": supplier.name}
        return {}

    def _org_payload(self, organization_type: str, organization_id: int) -> dict[str, Any]:
        if organization_type == "client":
            row = self.db.get(erp.Client, organization_id)
            return {"organization_type": "client", "client_id": organization_id, "name": row.name if row else None}
        if organization_type == "manufacturer":
            row = self.db.get(Manufacturer, organization_id)
            return {"organization_type": "manufacturer", "manufacturer_id": organization_id, "name": row.name if row else None}
        if organization_type == "supplier":
            row = self.db.get(Supplier, organization_id)
            return {"organization_type": "supplier", "supplier_id": organization_id, "name": row.name if row else None}
        return {"organization_type": organization_type, "name": None}

    def _domain_matches_name(self, domain: str, name: str) -> bool:
        stem = domain.split(".", 1)[0].replace("-", "").replace("_", "")
        normalized_name = re.sub(r"[^a-z0-9]", "", name.casefold())
        return len(stem) >= 4 and (stem in normalized_name or normalized_name in stem)

    def find_existing(self, normalized_email: str) -> erp.Contact | None:
        return self.db.query(erp.Contact).filter(or_(func.lower(erp.Contact.normalized_email) == normalized_email, func.lower(erp.Contact.email) == normalized_email)).first()

    def apply_decision(self, data: dict[str, Any], decision: dict[str, Any]) -> None:
        if not decision:
            return
        if decision.get("organization_type"):
            data["organization_type"] = decision["organization_type"]
        for field in ("client_id", "manufacturer_id", "supplier_id", "client_site_id"):
            if decision.get(field) is not None:
                data[field] = decision[field]
        if decision.get("organization_type") == "personal":
            data["client_id"] = data["manufacturer_id"] = data["supplier_id"] = None
        if decision.get("data_quality_status"):
            data["data_quality_status"] = decision["data_quality_status"]

    def issue(self, row_number: int, field: str, raw_value: Any, code: str, message: str, severity: str) -> dict[str, Any]:
        return {"row_number": row_number, "field_name": field, "raw_value": str(raw_value or ""), "error_code": code, "error_message": message, "severity": severity}

    def _audit(self, event_type: str, entity_type: str, entity_id: str | None, values: dict[str, Any]) -> None:
        safe_values = json.loads(json.dumps(values, default=str))
        self.db.add(AuditEvent(event_type=event_type, entity_type=entity_type, entity_id=entity_id, user_id=self.user_id, source=SOURCE, new_values=safe_values))


class CustomerContactsLargeImportService:
    def __init__(self, db: Session, user_id: int | None = None):
        self.db = db
        self.user_id = user_id

    def create_upload_job(self, content: bytes, filename: str, idempotency_key: str | None = None, allow_duplicate: bool = False) -> dict[str, Any]:
        suffix = Path(filename or "").suffix.casefold()
        if suffix not in ALLOWED_CONTACT_IMPORT_EXTENSIONS:
            raise ValueError("Upload must be .xlsx or .csv")
        if len(content) > CONTACT_IMPORT_MAX_FILE_SIZE_MB * 1024 * 1024:
            raise ValueError(f"Upload exceeds {CONTACT_IMPORT_MAX_FILE_SIZE_MB} MB")
        checksum = hashlib.sha256(content).hexdigest()
        previous = self.db.query(erp.CustomerContactImportJob).filter_by(file_checksum=checksum).order_by(erp.CustomerContactImportJob.id.desc()).first()
        if previous and not allow_duplicate:
            return {"duplicate_of_import_id": previous.id, "status": "duplicate_detected", "message": "This file appears to have been uploaded before."}
        stored_name = f"{datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')}-{uuid.uuid4().hex[:12]}-{safe_filename(filename)}{suffix}"
        path = CONTACT_IMPORT_STORAGE / stored_name
        path.write_bytes(content)
        batch = ImportBatch(source_type="customer_contacts", source_filename=filename, source_checksum=checksum, imported_by_id=self.user_id, status="uploaded", total_rows=0, notes=json.dumps({"stored_file_path": str(path)}))
        self.db.add(batch)
        self.db.flush()
        job = erp.CustomerContactImportJob(
            file_name=filename or stored_name,
            stored_file_path=str(path),
            file_size_bytes=len(content),
            file_checksum=checksum,
            status="queued",
            phase="file_validation",
            batch_size=CONTACT_IMPORT_BATCH_SIZE,
            created_by=self.user_id,
            import_batch_id=batch.id,
            idempotency_key=idempotency_key,
        )
        self.db.add(job)
        self._audit("customer_contacts_workbook_uploaded", "customer_contact_import_job", None, {"filename": filename, "checksum": checksum, "size": len(content)})
        self.db.commit()
        self.db.refresh(job)
        return self.job_payload(job)

    def create_upload_job_from_stored_file(self, path: Path, filename: str, file_size: int, checksum: str, idempotency_key: str | None = None, allow_duplicate: bool = False) -> dict[str, Any]:
        suffix = Path(filename or path.name).suffix.casefold()
        if suffix not in ALLOWED_CONTACT_IMPORT_EXTENSIONS:
            raise ValueError("Upload must be .xlsx or .csv")
        if file_size > CONTACT_IMPORT_MAX_FILE_SIZE_MB * 1024 * 1024:
            raise ValueError(f"Upload exceeds {CONTACT_IMPORT_MAX_FILE_SIZE_MB} MB")
        previous = self.db.query(erp.CustomerContactImportJob).filter_by(file_checksum=checksum).order_by(erp.CustomerContactImportJob.id.desc()).first()
        if previous and not allow_duplicate:
            path.unlink(missing_ok=True)
            return {"duplicate_of_import_id": previous.id, "status": "duplicate_detected", "message": "This file appears to have been uploaded before."}
        batch = ImportBatch(source_type="customer_contacts", source_filename=filename, source_checksum=checksum, imported_by_id=self.user_id, status="uploaded", total_rows=0, notes=json.dumps({"stored_file_path": str(path)}))
        self.db.add(batch)
        self.db.flush()
        job = erp.CustomerContactImportJob(
            file_name=filename or path.name,
            stored_file_path=str(path),
            file_size_bytes=file_size,
            file_checksum=checksum,
            status="queued",
            phase="file_validation",
            batch_size=CONTACT_IMPORT_BATCH_SIZE,
            created_by=self.user_id,
            import_batch_id=batch.id,
            idempotency_key=idempotency_key,
        )
        self.db.add(job)
        self._audit("customer_contacts_workbook_uploaded", "customer_contact_import_job", None, {"filename": filename, "checksum": checksum, "size": file_size})
        self.db.commit()
        self.db.refresh(job)
        return self.job_payload(job)

    def analyse_job(self, import_id: int) -> None:
        job = self._job(import_id)
        if job.status == "cancellation_requested":
            job.status = "cancelled"
            job.cancelled_at = datetime.now(UTC)
            self.db.commit()
            return
        if job.status not in {"queued", "uploaded", "failed"}:
            return
        job.status = "analysing"
        job.phase = "header_detection"
        job.started_at = job.started_at or datetime.now(UTC)
        self.db.commit()
        try:
            self._stage_file(job)
            job = self._job(import_id)
            if job.cancel_requested_at:
                job.status = "cancelled"
                job.phase = "cleanup"
                job.cancelled_at = datetime.now(UTC)
            else:
                job.status = "awaiting_review"
                job.phase = "preview_ready"
                job.progress_percent = 100
            self._sync_import_batch(job)
            self._audit("customer_contacts_preview_generated", "customer_contact_import_job", str(job.id), self.job_payload(job))
            self.db.commit()
        except Exception as exc:
            self.db.rollback()
            failed = self._job(import_id)
            failed.status = "failed"
            failed.phase = "parsing"
            failed.failed_at = datetime.now(UTC)
            failed.failure_message = str(exc)[:1000]
            self.db.commit()
            raise

    def confirm_job(self, import_id: int, idempotency_key: str | None = None) -> dict[str, Any]:
        job = self._job(import_id)
        if idempotency_key and job.idempotency_key == f"confirm:{idempotency_key}" and job.status in {"queued", "importing", "completed", "completed_with_errors"}:
            return self.job_payload(job)
        if job.status not in {"awaiting_review", "completed_with_errors", "cancelled"}:
            raise ValueError("Import must be awaiting review before confirmation")
        job.status = "queued"
        job.phase = "database_import"
        job.last_processed_row = 0
        job.idempotency_key = f"confirm:{idempotency_key}" if idempotency_key else job.idempotency_key
        self.db.commit()
        return self.job_payload(job)

    def run_import_job(self, import_id: int) -> None:
        job = self._job(import_id)
        if job.status not in {"queued", "importing"} or job.phase != "database_import":
            return
        job.status = "importing"
        job.started_at = job.started_at or datetime.now(UTC)
        self.db.commit()
        while True:
            job = self._job(import_id)
            if job.cancel_requested_at:
                job.status = "cancelled"
                job.cancelled_at = datetime.now(UTC)
                self.db.commit()
                return
            rows = (
                self.db.query(erp.CustomerContactImportStaging)
                .filter(erp.CustomerContactImportStaging.import_job_id == import_id)
                .filter(erp.CustomerContactImportStaging.validation_status != "error")
                .filter(erp.CustomerContactImportStaging.proposed_action.in_(["create_contact", "update_existing_contact"]))
                .filter(erp.CustomerContactImportStaging.decision_status.in_(["pending", "approved"]))
                .filter(erp.CustomerContactImportStaging.id > job.last_processed_row)
                .order_by(erp.CustomerContactImportStaging.id)
                .limit(CONTACT_IMPORT_DB_BATCH_SIZE)
                .all()
            )
            if not rows:
                job.status = "completed_with_errors" if job.error_rows else "completed"
                job.phase = "report_generation"
                job.completed_at = datetime.now(UTC)
                self._generate_report(job)
                self._sync_import_batch(job)
                self.db.commit()
                return
            try:
                self._import_stage_batch(job, rows)
                self.db.commit()
            except Exception as exc:
                self.db.rollback()
                job = self._job(import_id)
                job.status = "failed"
                job.failed_at = datetime.now(UTC)
                job.failure_message = "DATABASE_WRITE_FAILURE"
                self.db.add(AuditEvent(event_type="customer_contacts_import_failed", entity_type="customer_contact_import_job", entity_id=str(job.id), user_id=self.user_id, source=SOURCE, new_values={"error": str(exc)[:500]}))
                self.db.commit()
                return

    def rows(self, import_id: int, page: int = 1, page_size: int = 50, status: str = "", action: str = "", domain: str = "", organization_id: int | None = None, search: str = "", sort_by: str = "source_row_number", sort_direction: str = "asc") -> dict[str, Any]:
        page_size = min(max(page_size, 1), 200)
        query = self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id)
        if status:
            query = query.filter(erp.CustomerContactImportStaging.validation_status == status)
        if action:
            query = query.filter(erp.CustomerContactImportStaging.proposed_action == action)
        if domain:
            query = query.filter(erp.CustomerContactImportStaging.normalized_domain == domain.casefold())
        if organization_id:
            query = query.filter(erp.CustomerContactImportStaging.suggested_organization_id == organization_id)
        if search:
            like = f"%{search.casefold()}%"
            query = query.filter(or_(func.lower(erp.CustomerContactImportStaging.original_name).like(like), func.lower(erp.CustomerContactImportStaging.normalized_email).like(like), func.lower(erp.CustomerContactImportStaging.normalized_domain).like(like), func.lower(erp.CustomerContactImportStaging.suggested_organization_name).like(like)))
        total = query.count()
        sort_column = getattr(erp.CustomerContactImportStaging, sort_by, erp.CustomerContactImportStaging.source_row_number)
        query = query.order_by(sort_column.desc() if sort_direction == "desc" else sort_column.asc())
        items = [self.stage_payload(row) for row in query.offset((page - 1) * page_size).limit(page_size).all()]
        return {"items": items, "total": total, "page": page, "page_size": page_size, "summary": self.summary(import_id)}

    def bulk_decision(self, import_id: int, payload: dict[str, Any]) -> dict[str, int]:
        query = self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id)
        filters = payload.get("filters") or {}
        if filters.get("domain"):
            query = query.filter(erp.CustomerContactImportStaging.normalized_domain == str(filters["domain"]).casefold())
        if filters.get("automated_addresses"):
            query = query.filter(erp.CustomerContactImportStaging.is_automated_address.is_(True))
        if filters.get("shared_inboxes"):
            query = query.filter(erp.CustomerContactImportStaging.is_shared_inbox.is_(True))
        if filters.get("warning_code"):
            query = query.filter(erp.CustomerContactImportStaging.warning_code == filters["warning_code"])
        if filters.get("min_engagement") is not None:
            query = query.filter(erp.CustomerContactImportStaging.email_count >= int(filters["min_engagement"]))
        values = {"decision_status": "approved"}
        action = payload.get("action")
        if action:
            values["decision_action"] = action
            if action in {"reject", "ignore", "mark_personal"}:
                values["proposed_action"] = action
        if payload.get("organization_type") and payload.get("organization_id"):
            values["suggested_organization_type"] = payload["organization_type"]
            values["suggested_organization_id"] = int(payload["organization_id"])
            values["proposed_action"] = "update_existing_contact"
        updated = query.update(values, synchronize_session=False)
        self.db.add(AuditEvent(event_type="customer_contacts_bulk_decision", entity_type="customer_contact_import_job", entity_id=str(import_id), user_id=self.user_id, source=SOURCE, new_values={"updated": updated, **payload}))
        self.db.commit()
        return {"updated_rows": updated}

    def cancel(self, import_id: int) -> dict[str, Any]:
        job = self._job(import_id)
        job.status = "cancellation_requested"
        job.cancel_requested_at = datetime.now(UTC)
        self.db.commit()
        return self.job_payload(job)

    def resume(self, import_id: int) -> dict[str, Any]:
        job = self._job(import_id)
        if job.status not in {"failed", "cancelled", "completed_with_errors"}:
            raise ValueError("Only failed, cancelled, or partially completed jobs can be resumed")
        job.status = "queued"
        job.cancel_requested_at = None
        job.cancelled_at = None
        job.retry_attempts += 1
        self.db.commit()
        return self.job_payload(job)

    def summary(self, import_id: int) -> dict[str, Any]:
        job = self._job(import_id)
        return self.job_payload(job) | {
            "valid": self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id, validation_status="valid").count(),
            "warnings": self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id, validation_status="warning").count(),
            "errors": self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id, validation_status="error").count(),
            "duplicates": self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id).filter(erp.CustomerContactImportStaging.warning_code.like("%DUPLICATE%")).count(),
            "unmatched_domains": self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id, warning_code="UNMATCHED_ORGANIZATION").count(),
        }

    def report_csv(self, import_id: int) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["row", "status", "action", "name", "email", "domain", "emails_exchanged", "engagement", "phone", "role", "suggested_organization", "error_code", "warning_code", "decision"])
        for row in self.db.query(erp.CustomerContactImportStaging).filter_by(import_job_id=import_id).order_by(erp.CustomerContactImportStaging.source_row_number).yield_per(1000):
            writer.writerow([row.source_row_number, row.validation_status, row.proposed_action, row.original_name, row.normalized_email, row.normalized_domain, row.email_count, row.engagement_level, row.phone, row.role, row.suggested_organization_name, row.error_code, row.warning_code, row.decision_status])
        return output.getvalue().encode("utf-8")

    def _stage_file(self, job: erp.CustomerContactImportJob) -> None:
        path = Path(job.stored_file_path)
        rows_iter, closer = self._iter_source_rows(path)
        try:
            batch = []
            for row_number, raw in rows_iter:
                if row_number > CONTACT_IMPORT_MAX_ROWS + 1:
                    raise ValueError(f"Import exceeds CONTACT_IMPORT_MAX_ROWS={CONTACT_IMPORT_MAX_ROWS}")
                batch.append((row_number, raw))
                if len(batch) >= job.batch_size:
                    self._stage_batch(job, batch)
                    batch = []
            if batch:
                self._stage_batch(job, batch)
        finally:
            closer()

    def _stage_batch(self, job: erp.CustomerContactImportJob, batch: list[tuple[int, dict[str, Any]]]) -> None:
        domains = set()
        emails = set()
        normalized_rows = []
        seen_in_batch: dict[str, int] = {}
        for row_number, raw in batch:
            normalized = self._normalize_stage_row(row_number, raw)
            email = normalized.get("normalized_email")
            if email:
                if email in seen_in_batch:
                    normalized["validation_status"] = "warning"
                    normalized["proposed_action"] = "ignore"
                    normalized["warning_code"] = ",".join(filter(None, [normalized.get("warning_code"), "DUPLICATE_IN_FILE"]))
                    normalized["duplicate_of_row"] = seen_in_batch[email]
                    job.duplicate_rows += 1
                else:
                    seen_in_batch[email] = row_number
            normalized_rows.append(normalized)
            if email:
                emails.add(email)
            if normalized.get("normalized_domain"):
                domains.add(normalized["normalized_domain"])
        existing = {row.normalized_email or row.email.casefold(): row for row in self.db.query(erp.Contact).filter(or_(erp.Contact.normalized_email.in_(emails), func.lower(erp.Contact.email).in_(emails))).all()} if emails else {}
        staged = {
            row.normalized_email: row
            for row in self.db.query(erp.CustomerContactImportStaging)
            .filter(erp.CustomerContactImportStaging.import_job_id == job.id, erp.CustomerContactImportStaging.normalized_email.in_(emails))
            .all()
        } if emails else {}
        domain_matches = self._domain_matches(domains)
        for normalized in normalized_rows:
            email = normalized.get("normalized_email")
            if normalized.get("duplicate_of_row"):
                continue
            if email and email in staged:
                previous = staged[email]
                previous.warning_code = ",".join(filter(None, [previous.warning_code, "DUPLICATE_IN_FILE"]))
                previous.duplicate_row_numbers = ",".join(filter(None, [previous.duplicate_row_numbers, str(normalized.get("source_row_number"))]))
                previous.email_count = max(previous.email_count or 0, normalized.get("email_count") or 0)
                job.duplicate_rows += 1
                continue
            self._apply_batch_matches(normalized, existing, domain_matches)
            stage = erp.CustomerContactImportStaging(import_job_id=job.id, **{k: v for k, v in normalized.items() if hasattr(erp.CustomerContactImportStaging, k)})
            self.db.add(stage)
            self.db.flush()
            self._count_stage(job, stage)
        job.processed_rows += len(batch)
        job.total_rows = max(job.total_rows, job.processed_rows)
        job.last_processed_row = max(row_number for row_number, _ in batch)
        job.current_batch += 1
        job.last_completed_batch = job.current_batch
        job.progress_percent = min(99, int((job.processed_rows / max(job.total_rows, job.processed_rows, 1)) * 100))
        self._sync_import_batch(job)
        self.db.commit()

    def _normalize_stage_row(self, row_number: int, raw: dict[str, Any]) -> dict[str, Any]:
        name = str(raw.get("Name") or "").strip()
        email = normalize_email(raw.get("Email"))
        domain = email_domain(email)
        supplied_domain = str(raw.get("Company / Domain") or "").strip().casefold()
        email_count = parse_count(raw.get("Emails Exchanged"))
        phone, phone_original = normalize_phone(raw.get("Phone"))
        shared, contact_type = detect_shared(name, email)
        automated = detect_automated(name, email)
        warning_codes = []
        error_code = None
        proposed_action = "create_contact"
        validation_status = "valid"
        if not email:
            proposed_action = "ignore"
            validation_status = "skipped"
            warning_codes.append("MISSING_EMAIL")
        elif not EMAIL_RE.match(email):
            validation_status = "error"
            error_code = "INVALID_EMAIL"
        if supplied_domain and domain and supplied_domain != domain and "@" not in supplied_domain:
            warning_codes.append("DOMAIN_MISMATCH")
        if automated:
            warning_codes.append("AUTOMATED_ADDRESS")
        if shared:
            warning_codes.append("SHARED_INBOX_REVIEW")
        if phone_original and phone and len(re.sub(r"\D", "", phone)) < 5:
            warning_codes.append("INVALID_PHONE")
        if warning_codes and validation_status == "valid":
            validation_status = "warning"
        return {
            "source_row_number": row_number,
            "original_name": name,
            "original_email": str(raw.get("Email") or "").strip(),
            "normalized_email": email or None,
            "original_domain": supplied_domain or None,
            "normalized_domain": domain or supplied_domain or None,
            "email_count": email_count,
            "phone": phone,
            "role": str(raw.get("Role") or "").strip() or None,
            "notes": str(raw.get("Notes") or "").strip() or None,
            "proposed_action": proposed_action,
            "validation_status": validation_status,
            "error_code": error_code,
            "warning_code": ",".join(warning_codes) or None,
            "is_shared_inbox": shared,
            "is_automated_address": automated,
            "engagement_level": engagement_level(email_count),
            "contact_type": "shared_inbox" if automated else contact_type,
            "data_quality_status": "needs_review" if warning_codes else "needs_role",
            "validation_metadata": json.dumps({"raw": raw, "phone_original": phone_original}, default=str),
        }

    def _apply_batch_matches(self, normalized: dict[str, Any], existing: dict[str, erp.Contact], domain_matches: dict[str, dict[str, Any]]) -> None:
        email = normalized.get("normalized_email")
        domain = normalized.get("normalized_domain")
        if email and email in existing:
            normalized["existing_contact_id"] = existing[email].id
            normalized["proposed_action"] = "update_existing_contact"
            normalized["warning_code"] = ",".join(filter(None, [normalized.get("warning_code"), "DUPLICATE_EXISTING_CONTACT"]))
        if domain in PUBLIC_DOMAINS:
            normalized["suggested_organization_type"] = "personal"
            normalized["data_quality_status"] = "needs_review"
            normalized["warning_code"] = ",".join(filter(None, [normalized.get("warning_code"), "UNMATCHED_ORGANIZATION"]))
            return
        match = domain_matches.get(domain or "")
        if match:
            normalized["suggested_organization_type"] = match["organization_type"]
            normalized["suggested_organization_id"] = match["organization_id"]
            normalized["suggested_organization_name"] = match["name"]
        elif email:
            normalized["warning_code"] = ",".join(filter(None, [normalized.get("warning_code"), "UNMATCHED_ORGANIZATION"]))
            normalized["data_quality_status"] = "needs_organization_match"
            if normalized["validation_status"] == "valid":
                normalized["validation_status"] = "warning"

    def _domain_matches(self, domains: set[str]) -> dict[str, dict[str, Any]]:
        matches = {}
        if not domains:
            return matches
        mappings = self.db.query(erp.OrganizationDomainMapping).filter(erp.OrganizationDomainMapping.domain.in_(domains), erp.OrganizationDomainMapping.status == "approved").all()
        for mapping in mappings:
            matches[mapping.domain] = {"organization_type": mapping.organization_type, "organization_id": mapping.organization_id, "name": None}
        return matches

    def _import_stage_batch(self, job: erp.CustomerContactImportJob, rows: list[erp.CustomerContactImportStaging]) -> None:
        emails = [row.normalized_email for row in rows if row.normalized_email]
        existing = {row.normalized_email or row.email.casefold(): row for row in self.db.query(erp.Contact).filter(or_(erp.Contact.normalized_email.in_(emails), func.lower(erp.Contact.email).in_(emails))).all()} if emails else {}
        for row in rows:
            if row.proposed_action in {"reject", "ignore", "mark_personal"}:
                row.decision_status = "applied"
                if row.proposed_action == "reject":
                    job.rejected_rows += 1
                continue
            data = self._stage_to_contact_data(row)
            contact = existing.get(row.normalized_email)
            was_created = contact is None
            if was_created:
                contact = erp.Contact(name=row.original_name or row.normalized_email, display_name=row.original_name, email=row.normalized_email, normalized_email=row.normalized_email)
                self.db.add(contact)
            self._apply_contact_data(contact, data)
            self.db.flush()
            row.existing_contact_id = contact.id
            row.decision_status = "applied"
            job.created_rows += 1 if was_created else 0
            job.updated_rows += 0 if was_created else 1
            self.db.add(AuditEvent(event_type="customer_contact_created" if was_created else "customer_contact_updated", entity_type="contact", entity_id=str(contact.id), user_id=self.user_id, source=SOURCE, new_values={"import_job_id": job.id, "row": row.source_row_number, "email": contact.email}))
            job.last_processed_row = row.id
        job.progress_percent = min(99, job.progress_percent)

    def _stage_to_contact_data(self, row: erp.CustomerContactImportStaging) -> dict[str, Any]:
        org_type = row.suggested_organization_type
        return {
            "display_name": row.original_name,
            "email": row.normalized_email,
            "normalized_email": row.normalized_email,
            "email_domain": row.normalized_domain,
            "emails_exchanged": row.email_count,
            "engagement_level": row.engagement_level,
            "phone": row.phone,
            "role_title": row.role,
            "contact_type": row.contact_type,
            "organization_type": org_type or "unknown",
            "client_id": row.suggested_organization_id if org_type == "client" else None,
            "manufacturer_id": row.suggested_organization_id if org_type == "manufacturer" else None,
            "supplier_id": row.suggested_organization_id if org_type == "supplier" else None,
            "is_shared_inbox": row.is_shared_inbox,
            "is_automated_address": row.is_automated_address,
            "data_quality_status": row.data_quality_status,
        }

    def _apply_contact_data(self, contact: erp.Contact, data: dict[str, Any]) -> None:
        contact.display_name = contact.display_name or data.get("display_name")
        contact.name = contact.name or data.get("display_name") or data["email"]
        contact.email = data["email"]
        contact.normalized_email = data["normalized_email"]
        contact.email_domain = data.get("email_domain")
        contact.role_title = contact.role_title or data.get("role_title")
        if data.get("phone") and (not contact.phone or not contact.phone_verified):
            contact.phone = data["phone"]
        if int(data.get("emails_exchanged") or 0) > int(contact.emails_exchanged or 0):
            contact.emails_exchanged = data["emails_exchanged"]
            contact.engagement_level = data["engagement_level"]
        contact.contact_type = data.get("contact_type") or contact.contact_type
        contact.organization_type = data.get("organization_type") or contact.organization_type
        for field in ("client_id", "manufacturer_id", "supplier_id"):
            if getattr(contact, field, None) is None and data.get(field):
                setattr(contact, field, data[field])
        contact.is_shared_inbox = bool(data.get("is_shared_inbox"))
        contact.is_automated_address = bool(data.get("is_automated_address"))
        contact.data_quality_status = data.get("data_quality_status") or contact.data_quality_status
        contact.source = SOURCE
        contact.last_imported_at = datetime.now(UTC)

    def _iter_source_rows(self, path: Path):
        if path.suffix.casefold() == ".csv":
            file = path.open("r", encoding="utf-8-sig", newline="")
            sample = file.read(4096)
            file.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
            reader = csv.DictReader(file, dialect=dialect)
            def iterator():
                for index, row in enumerate(reader, start=2):
                    yield index, {column: row.get(column, "") for column in EXPECTED_COLUMNS}
            return iterator(), file.close
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheet = workbook[EXPECTED_SHEET] if EXPECTED_SHEET in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        header_map = {header: index for index, header in enumerate(headers)}
        def iterator():
            for index, values in enumerate(rows, start=2):
                yield index, {column: values[header_map[column]] if column in header_map and header_map[column] < len(values) else "" for column in EXPECTED_COLUMNS}
        return iterator(), workbook.close

    def _count_stage(self, job: erp.CustomerContactImportJob, stage: erp.CustomerContactImportStaging) -> None:
        if stage.validation_status == "error":
            job.error_rows += 1
        elif stage.validation_status == "skipped":
            job.skipped_rows += 1
        elif stage.validation_status == "warning":
            job.warning_rows += 1
        else:
            job.valid_rows += 1

    def _generate_report(self, job: erp.CustomerContactImportJob) -> None:
        report_path = CONTACT_IMPORT_STORAGE / f"customer-contact-import-{job.id}-report.csv"
        report_path.write_bytes(self.report_csv(job.id))
        job.report_path = str(report_path)

    def _sync_import_batch(self, job: erp.CustomerContactImportJob) -> None:
        if not job.import_batch_id:
            return
        batch = self.db.get(ImportBatch, job.import_batch_id)
        if not batch:
            return
        batch.status = job.status
        batch.total_rows = job.total_rows
        batch.processed_rows = job.processed_rows
        batch.successful_rows = job.created_rows + job.updated_rows
        batch.failed_rows = job.error_rows + job.rejected_rows
        if job.completed_at:
            batch.completed_at = job.completed_at

    def stage_payload(self, row: erp.CustomerContactImportStaging) -> dict[str, Any]:
        return {column.name: getattr(row, column.name) for column in row.__table__.columns}

    def job_payload(self, job: erp.CustomerContactImportJob) -> dict[str, Any]:
        elapsed = None
        if job.started_at:
            end = job.completed_at or job.failed_at or job.cancelled_at or datetime.now(UTC)
            start = job.started_at.replace(tzinfo=UTC) if job.started_at.tzinfo is None else job.started_at
            end = end.replace(tzinfo=UTC) if end.tzinfo is None else end
            elapsed = max(0, int((end - start).total_seconds()))
        return {column.name: getattr(job, column.name) for column in job.__table__.columns} | {"elapsed_seconds": elapsed}

    def _job(self, import_id: int) -> erp.CustomerContactImportJob:
        job = self.db.get(erp.CustomerContactImportJob, import_id)
        if not job:
            raise LookupError("Customer contacts import job not found")
        return job

    def _audit(self, event_type: str, entity_type: str, entity_id: str | None, values: dict[str, Any]) -> None:
        safe_values = json.loads(json.dumps(values, default=str))
        self.db.add(AuditEvent(event_type=event_type, entity_type=entity_type, entity_id=entity_id, user_id=self.user_id, source=SOURCE, new_values=safe_values))


def analyse_customer_contacts_job(import_id: int) -> None:
    with SessionLocal() as db:
        CustomerContactsLargeImportService(db).analyse_job(import_id)


def import_customer_contacts_job(import_id: int) -> None:
    with SessionLocal() as db:
        CustomerContactsLargeImportService(db).run_import_job(import_id)
