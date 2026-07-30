from __future__ import annotations

import hashlib
import io
import re
from datetime import UTC, date, datetime, time, timedelta
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app import erp_models as erp
from app.models.foundation import ImportBatch
from app.services.schedule_service import ScheduleService

DAY_NAMES = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday"]
TIME_RE = re.compile(r"\b(\d{1,2})(?::(\d{2}))?\s*(am|pm)?\b", re.I)
TABULAR_HEADERS = {
    "assigned_to",
    "engineer",
    "engineer_name",
    "task_name",
    "activity",
    "title",
    "due_date",
    "scheduled_date",
    "date",
    "hospital",
    "client",
    "status",
}


def parse_time(text: str) -> tuple[time | None, str | None]:
    match = TIME_RE.search(text or "")
    if not match:
        return None, "No confident time found."
    hour = int(match.group(1))
    minute = int(match.group(2) or 0)
    meridiem = (match.group(3) or "").lower()
    if meridiem == "pm" and hour < 12:
        hour += 12
    if meridiem == "am" and hour == 12:
        hour = 0
    if hour > 23 or minute > 59:
        return None, "Invalid time detected."
    return time(hour, minute), None


class ScheduleImportService:
    def __init__(self, db: Session, user_id: int | None = None, role: str = "viewer"):
        self.db = db
        self.user_id = user_id
        self.role = role

    def preview(self, content: bytes, filename: str) -> dict[str, Any]:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = self._find_sheet(workbook)
        week_of = self._find_week_of(sheet)
        headers = self._find_day_headers(sheet)
        warnings = []
        rows = self._preview_weekly_grid(sheet, week_of, headers)
        if not rows:
            rows = self._preview_tabular_sheet(sheet)
        if not week_of and not rows:
            warnings.append("Missing Week of date. Set dates before confirming import.")
        elif not week_of and headers:
            warnings.append("Missing Week of date. Set dates before confirming import.")
        return {
            "filename": filename,
            "checksum": hashlib.sha256(content).hexdigest(),
            "worksheet": sheet.title,
            "week_of": week_of.isoformat() if week_of else None,
            "warnings": warnings,
            "rows": rows,
        }

    def _preview_weekly_grid(self, sheet, week_of: date | None, headers: dict[int, str]) -> list[dict[str, Any]]:
        rows = []
        for column, day_name in headers.items():
            day_index = DAY_NAMES.index(day_name)
            event_date = week_of + timedelta(days=day_index) if week_of else None
            for row_index in range(1, sheet.max_row + 1):
                raw = sheet.cell(row=row_index, column=column).value
                if not raw or str(raw).strip().lower() == day_name:
                    continue
                for part in self._split_cell(str(raw)):
                    start_time, warning = parse_time(part)
                    start_datetime = datetime.combine(event_date, start_time) if event_date and start_time else None
                    end_datetime = start_datetime + timedelta(hours=1) if start_datetime else None
                    client = self._match_client(part)
                    rows.append({
                        "row_number": len(rows) + 1,
                        "day": day_name.title(),
                        "date": event_date.isoformat() if event_date else None,
                        "title": self._clean_title(part),
                        "source_reference": part,
                        "event_type": self._guess_type(part),
                        "status": "scheduled",
                        "priority": "normal",
                        "start_datetime": start_datetime.isoformat() if start_datetime else None,
                        "end_datetime": end_datetime.isoformat() if end_datetime else None,
                        "all_day": not bool(start_datetime),
                        "client_id": client.id if client else None,
                        "client_name": client.name if client else None,
                        "warnings": [warning] if warning else [],
                        "assignments": [],
                    })
        return rows

    def _preview_tabular_sheet(self, sheet) -> list[dict[str, Any]]:
        header_row = self._find_tabular_headers(sheet)
        if not header_row:
            return []
        row_number, headers = header_row
        rows = []
        for excel_row in range(row_number + 1, sheet.max_row + 1):
            values = {
                key: sheet.cell(row=excel_row, column=column).value
                for key, column in headers.items()
            }
            if not any(value not in (None, "") for value in values.values()):
                continue
            title = self._first_value(values, "task_name", "activity", "title") or "Imported schedule activity"
            scheduled_date = self._parse_date(self._first_value(values, "due_date", "scheduled_date", "date"))
            status = self._normalize_status(str(values.get("status") or "scheduled"))
            engineer = self._match_engineer(str(self._first_value(values, "assigned_to", "engineer", "engineer_name") or ""))
            client = self._match_client(str(self._first_value(values, "hospital", "client") or title))
            warnings = []
            if self._first_value(values, "assigned_to", "engineer", "engineer_name") and not engineer:
                warnings.append("Engineer was not found; event will import unassigned.")
            if not scheduled_date:
                warnings.append("No confident schedule date found.")
            start_datetime = datetime.combine(scheduled_date, time(9, 0)) if scheduled_date else None
            end_datetime = start_datetime + timedelta(hours=1) if start_datetime else None
            assignments = [{"engineer_id": engineer.id, "assignment_role": "lead"}] if engineer else []
            rows.append({
                "row_number": excel_row,
                "day": scheduled_date.strftime("%A") if scheduled_date else None,
                "date": scheduled_date.isoformat() if scheduled_date else None,
                "title": str(title).strip()[:255],
                "source_reference": self._source_reference(values),
                "event_type": self._guess_type(str(title)),
                "status": status,
                "priority": "normal",
                "start_datetime": start_datetime.isoformat() if start_datetime else None,
                "end_datetime": end_datetime.isoformat() if end_datetime else None,
                "all_day": not bool(start_datetime),
                "client_id": client.id if client else None,
                "client_name": client.name if client else None,
                "warnings": warnings,
                "assignments": assignments,
            })
        return rows

    def confirm(self, payload: dict[str, Any]) -> dict[str, Any]:
        checksum = payload.get("checksum")
        if checksum:
            duplicate = self.db.query(ImportBatch).filter_by(source_checksum=checksum, status="completed", source_type="engineer_schedule").first()
            if duplicate and not payload.get("allow_duplicate"):
                raise ValueError("This schedule file was already imported. Enable duplicate import to continue.")
        batch = ImportBatch(source_type="engineer_schedule", source_filename=payload.get("filename"), source_checksum=checksum, imported_by_id=self.user_id, status="processing", total_rows=len(payload.get("rows") or []))
        self.db.add(batch)
        self.db.flush()
        created = 0
        service = ScheduleService(self.db, self.user_id, self.role)
        for row in payload.get("rows") or []:
            if row.get("skip"):
                continue
            row["import_batch_id"] = batch.id
            row["source"] = "weekly_excel_import"
            service.create_event(dict(row))
            created += 1
        batch.status = "completed"
        batch.processed_rows = len(payload.get("rows") or [])
        batch.successful_rows = created
        batch.completed_at = datetime.now(UTC)
        self.db.commit()
        return {"import_batch_id": batch.id, "created": created}

    def _find_sheet(self, workbook):
        for sheet in workbook.worksheets:
            text = " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row[:8]).lower()
            if "weekly schedule" in text or "week of" in text:
                return sheet
        return workbook.active

    def _find_week_of(self, sheet) -> date | None:
        for row in sheet.iter_rows():
            for index, cell in enumerate(row):
                if str(cell.value or "").strip().lower().startswith("week of"):
                    for candidate in row[index + 1:index + 4]:
                        if isinstance(candidate.value, datetime):
                            return candidate.value.date()
                        if isinstance(candidate.value, date):
                            return candidate.value
                        try:
                            return datetime.fromisoformat(str(candidate.value)).date()
                        except Exception:
                            pass
        return None

    def _find_day_headers(self, sheet) -> dict[int, str]:
        headers = {}
        for row in sheet.iter_rows():
            for cell in row:
                label = str(cell.value or "").strip().lower()
                if label in DAY_NAMES:
                    headers[cell.column] = label
            if headers:
                return headers
        return {}

    def _find_tabular_headers(self, sheet) -> tuple[int, dict[str, int]] | None:
        for row in sheet.iter_rows():
            headers = {}
            for cell in row:
                key = self._normalize_header(cell.value)
                if key in TABULAR_HEADERS:
                    headers[key] = cell.column
            title_headers = {"task_name", "activity", "title"}
            date_headers = {"due_date", "scheduled_date", "date"}
            if title_headers & set(headers) and date_headers & set(headers):
                return row[0].row, headers
        return None

    def _normalize_header(self, value: Any) -> str:
        return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")

    def _split_cell(self, value: str) -> list[str]:
        parts = re.split(r"(?:\n|;|\s{2,})+", value)
        return [part.strip(" -\t") for part in parts if part.strip(" -\t")]

    def _clean_title(self, value: str) -> str:
        return re.sub(TIME_RE, "", value).strip(" -")[:255] or "Imported schedule activity"

    def _guess_type(self, value: str) -> str:
        text = value.lower()
        if "training" in text:
            return "training"
        if "install" in text:
            return "installation"
        if "contract" in text:
            return "contract_work"
        if "delivery" in text:
            return "delivery"
        if "pm" in text or "preventive" in text:
            return "preventive_maintenance"
        return "other"

    def _match_client(self, value: str):
        text = value.lower()
        if not text:
            return None
        clients = self.db.query(erp.Client).filter(erp.Client.name.isnot(None)).all()
        return next((client for client in clients if client.name.lower() in text), None)

    def _match_engineer(self, value: str):
        text = value.strip().lower()
        if not text:
            return None
        engineers = self.db.query(erp.Engineer).filter(erp.Engineer.engineer_name.isnot(None)).all()
        exact_match = next((engineer for engineer in engineers if engineer.engineer_name.lower() == text), None)
        return exact_match or next((engineer for engineer in engineers if engineer.engineer_name.lower() in text), None)

    def _first_value(self, values: dict[str, Any], *keys: str) -> Any:
        for key in keys:
            value = values.get(key)
            if value not in (None, ""):
                return value
        return None

    def _parse_date(self, value: Any) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if value in (None, ""):
            return None
        text = str(value).strip()
        for parser in (datetime.fromisoformat,):
            try:
                return parser(text).date()
            except ValueError:
                pass
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        return None

    def _normalize_status(self, value: str) -> str:
        text = value.strip().lower().replace(" ", "_")
        return text if text in {"draft", "scheduled", "confirmed", "in_progress", "completed", "cancelled", "postponed"} else "scheduled"

    def _source_reference(self, values: dict[str, Any]) -> str:
        parts = []
        for key, value in values.items():
            if value not in (None, ""):
                parts.append(f"{key}: {value}")
        return "; ".join(parts)
